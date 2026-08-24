"""Chart A / Table A workbook parsers for PDA Board.

Pure functions over openpyxl workbooks — no Spark, no pandas — so the same
module runs unchanged in local tests and inside a Databricks job.

Anchors (verified against the IN PDA July 2026 example files, 2026-07-18):

  Chart A ('County Summary' sheet)
    B3  "<County> (<population>) - $<county threshold>"   county + population + threshold
    D3  FEMA Inspector
    E3  Emergency Work total (A+B)   G3 Permanent Work total (C-G)   L3 Total
    B5  applicant table header: Applicant|Type|Status|A|B|C|D|E|F|G|Applicant Total
    B6+ applicant rows until first blank Applicant cell
  Chart A ('Comments' sheet): A=Applicant, B=Comments, rows 2+ (blank rows read 0)

  Table A — TWO template generations exist in the wild; both are parsed by
  label, not fixed cells (parse_table_a scans the Input Sheet header for
  'State Code:' / 'Population:' / 'Region:' / ... labels and reads the cell to
  the right of each):
    * v2 (IN PDA July 2026): has an 'RVAR' sheet. Input Sheet D1 = state code,
      D2 = state population. Per-county table on 'Summary' under a
      'County|Population' header.
    * legacy (MI/WI PDAs, May 2026): NO 'RVAR' sheet ('SummaryPage' +
      'PCI Indicator' instead). Input Sheet D1 = Region number, D2 = state
      code; state population is NOT on the Input Sheet — it comes from the
      'PCI Indicator' sheet (row under the 'STATE/TERRITORY' header, which
      also carries both PCIs + the incident name/dates block). Per-county
      table on 'SummaryPage' under a 'County|Subrecipient' header
      (C..I = cats, J = subtotal, K = county population).
  Both: subrecipient rows follow the 'Subrecipient|County|Status|A..G' header
  row (row 4), until a 'Total' row (v2) or end of data (legacy).
  Table A ('Reductions', hidden): County|Applicant|Cat|Original|Eligible|
    Reduction|Reduction Category|Notes, header on row 2 (same in both).

  Chart A format wrinkles (MI/WI May 2026 files): the comments sheet may be
  named 'Deduction Comments' instead of 'Comments' (same A=Applicant,
  B=Comments layout); filenames carry no incident-period parenthetical, so
  event_label is None and the event must come from the Table A / caller.

Trust rules (why some fields exist twice):
  * County identity comes from the B3 cell; the filename is only a fallback
    (real files named "Blank Chart A - ..." exist in the wild).
  * Every total is recomputed from the Cat A-G inputs; the workbook's own
    total cells are kept as ``wb_*`` for reconciliation and NEVER displayed.
  * Nothing is silently dropped: every anomaly lands in ``flags``.
"""
from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Optional

import openpyxl

warnings.filterwarnings("ignore", module="openpyxl")  # data-validation ext noise

CATS = ("A", "B", "C", "D", "E", "F", "G")
EMERGENCY_CATS = ("A", "B")  # Emergency Work; C-G = Permanent Work
TOL = 0.005  # dollar reconciliation tolerance (half a cent)
MAX_APPLICANT_ROWS = 500

COUNTY_RE = re.compile(
    r"^(?P<name>.+?)\s*\(\s*(?P<pop>[\d,]+)\s*\)\s*-\s*\$\s*(?P<thresh>[\d,.]+)\s*$"
)
# "(June 6 - 11, 2026)" / "(June 16 - 26, 2026)" / "(June 6 - July 2, 2026)"
EVENT_RE = re.compile(
    r"\((?P<label>[A-Za-z]+\s*\d{1,2}\s*-\s*(?:[A-Za-z]+\s*)?\d{1,2},\s*\d{4})\)"
)


# ---------------------------------------------------------------- helpers

def _num(v: Any) -> float:
    """Cell value -> float dollars; blanks and junk -> 0.0."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("$", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def _opt_num(v: Any) -> Optional[float]:
    """Like _num but preserves 'absent' (for wb_* reconciliation fields)."""
    if v is None or v == "":
        return None
    return _num(v)


def _clean(v: Any) -> str:
    return str(v).strip() if v not in (None, "") else ""


def _rich_html(v: Any) -> str:
    """Chart A comment cell -> minimal safe HTML preserving the workbook's own
    formatting (bold/italic/underline runs + alt-enter line breaks). Plain
    strings just get escaped + <br>. Pipeline-generated; rendered as-is by the
    board."""
    from html import escape
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
    except ImportError:                                    # very old openpyxl
        return escape(_clean(v)).replace("\n", "<br>")
    if not isinstance(v, CellRichText):
        return escape(_clean(v)).replace("\n", "<br>")
    parts = []
    for run in v:
        if isinstance(run, TextBlock):
            t = escape(run.text).replace("\n", "<br>")
            f = run.font
            if f is not None and f.u:
                t = f"<u>{t}</u>"
            if f is not None and f.i:
                t = f"<i>{t}</i>"
            if f is not None and f.b:
                t = f"<b>{t}</b>"
            parts.append(t)
        else:
            parts.append(escape(str(run)).replace("\n", "<br>"))
    return "".join(parts).strip()


def norm_status(raw: Any) -> str:
    s = _clean(raw).lower()
    if not s:
        return "blank"
    if s.startswith("complete") or s in ("validated", "done"):
        return "complete"
    if "progress" in s:
        return "in_progress"
    return "other"


def event_label_from_filename(filename: str) -> Optional[str]:
    m = EVENT_RE.search(filename)
    if not m:
        return None
    return normalize_event_label(m.group("label"))


def normalize_event_label(label: str) -> str:
    s = " ".join(str(label).split())
    s = re.sub(r"\s*-\s*", " - ", s)
    s = re.sub(r"\s*,\s*", ", ", s)
    return s


def event_key(label: Optional[str]) -> Optional[str]:
    """Whitespace/case-insensitive join key ('June 6 - 11, 2026' == 'June 6-11,2026')."""
    if not label:
        return None
    return re.sub(r"\s+", "", label).lower()


def county_short(name: Optional[str]) -> Optional[str]:
    """'Floyd County' -> 'Floyd' (join key against dim_county / folder names)."""
    if not name:
        return None
    return re.sub(r"\s+County$", "", name.strip(), flags=re.I).strip() or None


def county_from_filename(filename: str) -> Optional[str]:
    """Best-effort county from a Chart A filename. Fallback only — cells win.

    'Carroll County Chart A - Indiana (June 6 - 11, 2026).xlsx' -> 'Carroll'
    'Decatur - Indiana (June 6 - 11, 2026).xlsx'                -> 'Decatur'
    'Blank Chart A - Indiana (2026).xlsx'                       -> None
    'REMC Chart A - Indiana (...)'                              -> 'REMC'
    """
    s = Path(filename).stem
    s = re.sub(r"\(.*?\)", " ", s)
    s = re.sub(r"\bCompleted\b", " ", s, flags=re.I)
    s = s.split(" - ")[0]
    s = re.sub(r"\bChart\s*A\b.*$", " ", s, flags=re.I)
    s = re.sub(r"\bCounty\b", " ", s, flags=re.I)
    s = " ".join(s.split()).strip(" -_")
    if not s or s.lower() == "blank":
        return None
    return s


def is_blank_named(filename: str) -> bool:
    return Path(filename).stem.strip().lower().startswith("blank")


# ---------------------------------------------------------------- dataclasses

@dataclass
class Applicant:
    applicant: str
    applicant_type: str
    status_raw: str
    status_norm: str
    cats: dict  # {"A": float, ...}
    total_calc: float
    wb_row_total: Optional[float]
    comment: str = ""
    comment_html: str = ""

    def to_row(self) -> dict:
        d = {
            "applicant": self.applicant,
            "applicant_type": self.applicant_type,
            "status_raw": self.status_raw,
            "status_norm": self.status_norm,
        }
        d.update({f"cat_{c.lower()}": self.cats[c] for c in CATS})
        d["total_calc"] = self.total_calc
        d["wb_row_total"] = self.wb_row_total
        d["comment"] = self.comment
        d["comment_html"] = self.comment_html
        return d


@dataclass
class ChartA:
    source_file: str
    filename: str
    county_cell_raw: Optional[str]
    county: Optional[str]        # full name, e.g. "Floyd County"
    county_key: Optional[str]    # short join key, e.g. "Floyd"
    county_source: Optional[str]  # "cell" | "filename" | None
    entity_type: str             # "county" | "remc"
    population: Optional[int]
    county_threshold: Optional[float]
    inspector: Optional[str]
    event_label: Optional[str]
    event_key: Optional[str]
    applicants: list = field(default_factory=list)
    emergency_calc: float = 0.0
    permanent_calc: float = 0.0
    total_calc: float = 0.0
    wb_emergency: Optional[float] = None
    wb_permanent: Optional[float] = None
    wb_total: Optional[float] = None
    is_template: bool = False
    is_empty: bool = False
    flags: list = field(default_factory=list)

    @property
    def content_status(self) -> str:
        """What the cells themselves say (independent of folder stage)."""
        if self.is_empty:
            return "empty"
        if all(a.status_norm == "complete" for a in self.applicants):
            return "complete"
        return "active"

    def flag(self, code: str, detail: str = "") -> None:
        self.flags.append({"code": code, "detail": detail})

    def to_dict(self) -> dict:
        d = asdict(self)
        d["applicants"] = [a.to_row() for a in self.applicants]
        d["content_status"] = self.content_status
        return d


# ---------------------------------------------------------------- sniffing

def sniff_workbook(path) -> str:
    """'chart_a' | 'table_a' | 'other' from sheet names alone."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = set(wb.sheetnames)
        wb.close()
    except Exception:
        return "other"
    if "County Summary" in names:
        return "chart_a"
    # v2 has RVAR; the legacy (pre-RVAR) template has SummaryPage/PCI Indicator
    if "Input Sheet" in names and names & {"RVAR", "SummaryPage", "PCI Indicator"}:
        return "table_a"
    return "other"


# ---------------------------------------------------------------- Chart A

def parse_chart_a(path) -> ChartA:
    path = Path(path)
    # rich_text keeps the Comments cells' bold/italic runs (surfaced on the
    # board); _clean() str-coerces CellRichText for every other string cell.
    wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
    ws = wb["County Summary"]

    b3 = _clean(ws["B3"].value)
    m = COUNTY_RE.match(b3) if b3 else None
    fn_county = county_from_filename(path.name)

    if m:
        county = m.group("name").strip()
        county_source = "cell"
        population = int(m.group("pop").replace(",", ""))
        threshold = _num(m.group("thresh"))
    else:
        county = f"{fn_county} County" if fn_county and "remc" not in fn_county.lower() else fn_county
        county_source = "filename" if fn_county else None
        population = None
        threshold = None

    entity_type = "remc" if "remc" in (county or path.name).lower() else "county"

    ca = ChartA(
        source_file=str(path),
        filename=path.name,
        county_cell_raw=b3 or None,
        county=county,
        county_key=county_short(county) if entity_type == "county" else county,
        county_source=county_source,
        entity_type=entity_type,
        population=population,
        county_threshold=threshold,
        inspector=_clean(ws["D3"].value) or None,
        event_label=event_label_from_filename(path.name),
        event_key=event_key(event_label_from_filename(path.name)),
        wb_emergency=_opt_num(ws["E3"].value),
        wb_permanent=_opt_num(ws["G3"].value),
        wb_total=_opt_num(ws["L3"].value),
    )

    if b3 and not m:
        ca.flag("COUNTY_CELL_UNPARSEABLE", f"B3={b3!r}")
    if not b3:
        ca.flag("NO_COUNTY_CELL", "County Summary!B3 is blank (county dropdown not set)")

    # -- applicant rows: B6.. until first blank Applicant
    r = 6
    while r < 6 + MAX_APPLICANT_ROWS:
        name = ws.cell(row=r, column=2).value  # B
        if name in (None, "", 0):
            break
        cats = {c: _num(ws.cell(row=r, column=5 + i).value) for i, c in enumerate(CATS)}  # E..K
        total_calc = round(sum(cats.values()), 2)
        wb_row_total = _opt_num(ws.cell(row=r, column=12).value)  # L
        app = Applicant(
            applicant=_clean(name),
            applicant_type=_clean(ws.cell(row=r, column=3).value),   # C
            status_raw=_clean(ws.cell(row=r, column=4).value),        # D
            status_norm=norm_status(ws.cell(row=r, column=4).value),
            cats=cats,
            total_calc=total_calc,
            wb_row_total=wb_row_total,
        )
        if wb_row_total is not None and abs(wb_row_total - total_calc) > TOL:
            ca.flag("ROW_TOTAL_MISMATCH",
                    f"{app.applicant}: wb ${wb_row_total:,.2f} vs calc ${total_calc:,.2f}")
        ca.applicants.append(app)
        r += 1

    # -- comments joined by applicant name ('Comments' in the IN template,
    #    'Deduction Comments' in the MI/WI legacy template — same layout)
    comments_sheet = next((s for s in ("Comments", "Deduction Comments")
                           if s in wb.sheetnames), None)
    if comments_sheet:
        cws = wb[comments_sheet]
        comments = {}
        for rr in range(2, min(cws.max_row, 500) + 1):
            who = cws.cell(row=rr, column=1).value
            if who in (None, "", 0):
                continue
            raw = cws.cell(row=rr, column=2).value
            txt = _clean(raw)
            if txt:
                comments[_clean(who)] = (txt, _rich_html(raw))
        matched = set()
        for app in ca.applicants:
            if app.applicant in comments:
                app.comment, app.comment_html = comments[app.applicant]
                matched.add(app.applicant)
        for who in set(comments) - matched:
            ca.flag("ORPHAN_COMMENT", f"comment for {who!r} matches no applicant row")

    # -- recomputed totals (the displayed numbers)
    ca.emergency_calc = round(sum(a.cats[c] for a in ca.applicants for c in EMERGENCY_CATS), 2)
    ca.permanent_calc = round(sum(a.cats[c] for a in ca.applicants for c in CATS if c not in EMERGENCY_CATS), 2)
    ca.total_calc = round(ca.emergency_calc + ca.permanent_calc, 2)

    # -- reconciliation against the workbook's own cells
    for label, wb_v, calc in (("total L3", ca.wb_total, ca.total_calc),
                              ("emergency E3", ca.wb_emergency, ca.emergency_calc),
                              ("permanent G3", ca.wb_permanent, ca.permanent_calc)):
        if wb_v is not None and abs(wb_v - calc) > TOL:
            ca.flag("COUNTY_TOTAL_MISMATCH", f"{label}: wb ${wb_v:,.2f} vs calc ${calc:,.2f}")

    # -- template / empty classification
    ca.is_empty = not ca.applicants
    blank_named = is_blank_named(path.name)
    ca.is_template = blank_named and ca.is_empty and not m
    if blank_named and (ca.applicants or ca.total_calc > 0 or m):
        ca.flag("BLANK_NAMED_WITH_DATA",
                "file is named 'Blank Chart A' but carries real data — county taken from B3")
    if ca.is_empty and not ca.is_template:
        ca.flag("EMPTY_CHART", "no applicant rows")
    if county_source == "filename":
        ca.flag("COUNTY_FROM_FILENAME", f"county guessed from filename: {fn_county!r}")

    wb.close()
    return ca


# ---------------------------------------------------------------- Table A

@dataclass
class TableA:
    source_file: str
    filename: str
    meta: dict = field(default_factory=dict)
    counties: list = field(default_factory=list)       # per-county entered rows
    subrecipients: list = field(default_factory=list)  # Input Sheet grain
    reductions: list = field(default_factory=list)
    flags: list = field(default_factory=list)

    def flag(self, code: str, detail: str = "") -> None:
        self.flags.append({"code": code, "detail": detail})

    def to_dict(self) -> dict:
        return asdict(self)


def _input_labels(ws, max_row=3, max_col=13) -> dict:
    """Map 'Label:' cells in the Input Sheet header block -> the value cell to
    their right. Both Table A generations carry the same labels in different
    columns, so this is the format-proof way to read them."""
    out = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().endswith(":"):
                key = v.strip().rstrip(":").strip().lower()
                out.setdefault(key, ws.cell(row=r, column=c + 1).value)
    return out


def _find_header_row(ws, col_a, col_b, max_row=60) -> Optional[int]:
    for r in range(1, min(ws.max_row, max_row) + 1):
        if (_clean(ws.cell(row=r, column=1).value) == col_a
                and _clean(ws.cell(row=r, column=2).value) == col_b):
            return r
    return None


def parse_table_a(path) -> TableA:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ta = TableA(source_file=str(path), filename=path.name)
    fmt = "v2" if "RVAR" in wb.sheetnames else "legacy"

    ws = wb["Input Sheet"]
    lbl = _input_labels(ws)
    # 'Couties:' (sic, v2) / 'DATA ENTRY:' (legacy) also match — harmless.
    ta.meta = {
        "format": fmt,
        "state_code": _clean(lbl.get("state code")) or None,
        "region": _clean(lbl.get("region")) or None,
        "state_population": int(round(_num(lbl.get("population")))) or None,
        "state_pci": _opt_num(lbl.get("state per capita")),
        "county_pci": _opt_num(lbl.get("county per capita")),
        "event_type": _clean_meta_str(lbl.get("event type")),
        "pda_start": _fmt_date(lbl.get("pda start")),
        "incident_start": _fmt_date(lbl.get("incident start")),
        "incident_end": _fmt_date(lbl.get("incident end")),
        "incident_name": None,
        "state_name": None,
    }

    # legacy: state population (and a second chance at PCIs + incident meta)
    # lives on the 'PCI Indicator' sheet, not the Input Sheet. v2 keeps
    # everything on the Input Sheet — its PCI Indicator sheet can carry stale
    # template leftovers, so only the LEGACY format reads it.
    if fmt == "legacy" and "PCI Indicator" in wb.sheetnames:
        pws = wb["PCI Indicator"]
        for r in range(1, min(pws.max_row, 40) + 1):
            h = _clean(pws.cell(row=r, column=1).value).upper()
            if h == "STATE/TERRITORY":
                ta.meta["state_name"] = _clean(pws.cell(row=r + 1, column=1).value) or None
                ta.meta["state_population"] = (ta.meta["state_population"]
                    or int(round(_num(pws.cell(row=r + 1, column=2).value))) or None)
                ta.meta["state_pci"] = ta.meta["state_pci"] or _opt_num(
                    pws.cell(row=r + 1, column=3).value)
                ta.meta["county_pci"] = ta.meta["county_pci"] or _opt_num(
                    pws.cell(row=r + 1, column=4).value)
            elif h == "INCIDENT NAME":
                ta.meta["incident_name"] = _clean_meta_str(
                    pws.cell(row=r + 1, column=1).value)
                ta.meta["incident_start"] = ta.meta["incident_start"] or _fmt_date(
                    pws.cell(row=r + 1, column=3).value)
                ta.meta["incident_end"] = ta.meta["incident_end"] or _fmt_date(
                    pws.cell(row=r + 1, column=4).value)
                ta.meta["event_type"] = ta.meta["event_type"] or _clean_meta_str(
                    pws.cell(row=r + 1, column=5).value)

    if not ta.meta["state_code"]:
        ta.flag("NO_STATE_CODE", "no 'State Code:' label found on Input Sheet")

    # subrecipient rows: under the 'Subrecipient|County|...' header, until a
    # 'Total' row (v2) or end of data (legacy has no Total row)
    hdr = _find_header_row(ws, "Subrecipient", "County", max_row=10) or 4
    for r in range(hdr + 1, min(ws.max_row, 2000) + 1):
        a = _clean(ws.cell(row=r, column=1).value)
        if a.lower() == "total":
            break
        county = _clean(ws.cell(row=r, column=2).value)
        if not a and not county:
            continue
        cats = {c: _num(ws.cell(row=r, column=4 + i).value) for i, c in enumerate(CATS)}  # D..J
        ta.subrecipients.append({
            "subrecipient": a, "county": county,
            "status": _clean(ws.cell(row=r, column=3).value),
            **{f"cat_{c.lower()}": cats[c] for c in CATS},
            "total_calc": round(sum(cats.values()), 2),
        })

    # Per-county table.
    #   v2:     'Summary' sheet, 'County|Population' header, C..I cats, J total,
    #           K = county PCI target $, M = target label
    #   legacy: 'SummaryPage' sheet, 'County|Subrecipient' header, C..I cats,
    #           J subtotal, K = county population (target computed from pop x PCI)
    parsed_counties = False
    if "Summary" in wb.sheetnames:
        sws = wb["Summary"]
        hdr = _find_header_row(sws, "County", "Population")
        if hdr is not None:
            parsed_counties = True
            for r in range(hdr + 1, min(sws.max_row, hdr + 400) + 1):
                county = _clean(sws.cell(row=r, column=1).value)
                if not county:
                    break
                cats = {c: _num(sws.cell(row=r, column=3 + i).value) for i, c in enumerate(CATS)}  # C..I
                ta.counties.append({
                    "county": county,
                    "population": int(_num(sws.cell(row=r, column=2).value)) or None,
                    **{f"cat_{c.lower()}": cats[c] for c in CATS},
                    "total_entered": round(sum(cats.values()), 2),
                    "wb_total": _opt_num(sws.cell(row=r, column=10).value),        # J
                    "county_pci_target": _opt_num(sws.cell(row=r, column=11).value),  # K
                    "target_label": _clean(sws.cell(row=r, column=13).value) or None,  # M
                })
    if not parsed_counties and "SummaryPage" in wb.sheetnames:
        sws = wb["SummaryPage"]
        hdr = _find_header_row(sws, "County", "Subrecipient", max_row=80)
        if hdr is not None:
            parsed_counties = True
            cpci = ta.meta.get("county_pci")
            for r in range(hdr + 1, min(sws.max_row, hdr + 400) + 1):
                county = _clean(sws.cell(row=r, column=1).value)
                if not county:
                    # blank col A = an expanded per-subrecipient detail row
                    # under the county above (already counted in its county
                    # row) — skip, don't stop
                    continue
                if county.lower() in ("grand total", "total"):
                    break
                pop = int(_num(sws.cell(row=r, column=11).value)) or None  # K
                cats = {c: _num(sws.cell(row=r, column=3 + i).value) for i, c in enumerate(CATS)}  # C..I
                ta.counties.append({
                    "county": county,
                    "population": pop,
                    **{f"cat_{c.lower()}": cats[c] for c in CATS},
                    "total_entered": round(sum(cats.values()), 2),
                    "wb_total": _opt_num(sws.cell(row=r, column=10).value),   # J
                    "county_pci_target": round(pop * cpci, 2) if pop and cpci else None,
                    "target_label": None,
                })
    if not parsed_counties:
        ta.flag("SUMMARY_HEADER_NOT_FOUND",
                "no per-county header row found on Summary/SummaryPage")

    if "Reductions" in wb.sheetnames:
        rws = wb["Reductions"]
        for r in range(3, min(rws.max_row, 500) + 1):
            county = _clean(rws.cell(row=r, column=2).value)     # B
            applicant = _clean(rws.cell(row=r, column=3).value)  # C
            if not county and not applicant:
                continue
            ta.reductions.append({
                "county": county, "applicant": applicant,
                "cat": _clean(rws.cell(row=r, column=4).value),
                "original": _num(rws.cell(row=r, column=5).value),
                "eligible": _num(rws.cell(row=r, column=6).value),
                "reduction": _num(rws.cell(row=r, column=7).value),
                "reduction_category": _clean(rws.cell(row=r, column=8).value),
                "notes": _clean(rws.cell(row=r, column=9).value),
            })

    wb.close()
    return ta


def _fmt_date(v: Any) -> Optional[str]:
    if v in (None, "", 0):
        return None
    if hasattr(v, "strftime"):
        try:
            if getattr(v, "year", 2000) < 1950:   # 0-valued date cell (1900 epoch)
                return None
            return v.strftime("%Y-%m-%d")
        except ValueError:
            return None
    s = _clean(v)
    return s or None


def _clean_meta_str(v: Any) -> Optional[str]:
    """Meta text cell -> str, treating the template's fillers (0 / TBD /
    blank) as absent."""
    s = _clean(v)
    return None if s in ("", "0", "TBD") else s
