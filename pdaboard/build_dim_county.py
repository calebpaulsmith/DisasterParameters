"""Build a STATE's county dimension + map geometry for PDA Board.

Outputs (committed):
  data/dim_county_<st>.csv     fips,name,name_full,population,lat,lon
  data/<st>_counties_geo.json  GeoJSON (id = 5-digit FIPS)

Sources:
  * Populations: the hidden 'County Population and Threshold' sheet inside any
    of the PDA's Chart A workbooks — so the board's denominators match the
    PDA's own numbers exactly. Two template layouts exist (IN 2026: name col E
    / pop col F; MI/WI legacy: name col D / pop col E); the reader sniffs for
    whichever column pair holds ('<Name> County', number).
  * FIPS + geometry: a US counties GeoJSON (Census 20m, plotly datasets
    mirror) — counties are matched BY NAME within the state (normalized:
    case/space/period-insensitive), which is robust across states (no
    alphabetical-FIPS-rule assumption). Every county must match or the build
    fails loudly. Centroids = mean of the largest exterior ring.

Usage:
  python build_dim_county.py <state> [chart_a.xlsx] [counties_us.json]
  e.g. python build_dim_county.py mi "PDAExamples/_17-19 import/MI May 2026/.../Alcona....xlsx"
"""
from __future__ import annotations

import csv
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", module="openpyxl")

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
DATA = HERE / "data"

STATE_FIPS = {"in": "18", "mi": "26", "wi": "55", "il": "17", "mn": "27", "oh": "39"}
EXPECTED_N = {"in": 92, "mi": 83, "wi": 72, "il": 102, "mn": 87, "oh": 88}
# workbook-spelling -> census-spelling (normalized), for typos baked into the
# Chart A templates' hidden population sheet (real example: WI "Fond Du Loc")
NAME_ALIASES = {"fondduloc": "fonddulac", "lacross": "lacrosse"}


def _norm(name: str) -> str:
    return re.sub(r"[^a-z]", "", name.lower())


def read_populations_national(workbook_path: Path, state: str) -> list:
    """Fallback source: a Table A master's hidden 'County Populations' sheet —
    a NATIONAL lookup (state code | county | population), so any state's dim
    can be built from any Table A on hand (no Chart A for that state needed)."""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb["County Populations"]
    st = state.upper()
    out = []
    for r in ws.iter_rows(values_only=True):
        if len(r) < 4:
            continue
        if isinstance(r[1], str) and r[1].strip().upper() == st \
                and isinstance(r[2], str) and isinstance(r[3], (int, float)):
            full = r[2].strip()
            short = full[:-7].strip() if full.endswith(" County") else full
            out.append({"name": short, "name_full": f"{short} County",
                        "population": int(r[3])})
    wb.close()
    return out


def read_populations(chart_a_path: Path) -> list:
    wb = openpyxl.load_workbook(chart_a_path, data_only=True)
    ws = wb["County Population and Threshold"]
    grid = list(ws.iter_rows(values_only=True))
    wb.close()
    # sniff the (name, population) column pair: most rows where col c is
    # '<Name> County' text and col c+1 is a plain number
    best, best_n = None, 0
    ncols = max(len(r) for r in grid)
    for c in range(ncols - 1):
        n = sum(1 for r in grid
                if len(r) > c + 1 and isinstance(r[c], str)
                and r[c].strip().endswith(" County")
                and isinstance(r[c + 1], (int, float)))
        if n > best_n:
            best, best_n = c, n
    if best is None or best_n < 10:
        raise SystemExit(f"no county/population column pair found in {chart_a_path}")
    out = []
    for r in grid:
        if len(r) > best + 1 and isinstance(r[best], str) \
                and r[best].strip().endswith(" County") \
                and isinstance(r[best + 1], (int, float)):
            full = r[best].strip()
            out.append({"name": full[:-7].strip(), "name_full": full,
                        "population": int(r[best + 1])})
    return out


def largest_ring(geom: dict) -> list:
    if geom["type"] == "Polygon":
        return geom["coordinates"][0]
    rings = [poly[0] for poly in geom["coordinates"]]      # MultiPolygon
    return max(rings, key=len)


def build_geo(us_geojson: Path, counties: list, state_fips: str) -> dict:
    src = json.loads(us_geojson.read_text(encoding="utf-8"))
    feats = [f for f in src["features"]
             if f.get("properties", {}).get("STATE") == state_fips]
    by_name = {_norm(f["properties"]["NAME"]): f for f in feats}
    out_feats = []
    for c in counties:
        key = _norm(c["name"])
        key = NAME_ALIASES.get(key, key)
        f = by_name.get(key)
        if f is not None and key != _norm(c["name"]):
            # take the census spelling forward (board labels, join keys)
            print(f"  alias: workbook {c['name']!r} -> {f['properties']['NAME']!r}")
            c["name"] = f["properties"]["NAME"].strip()
            c["name_full"] = c["name"] + " County"
        if f is None:
            raise SystemExit(f"county {c['name']!r} not found in geojson for "
                             f"STATE={state_fips} — name variant? "
                             f"({sorted(by_name)[:5]}...)")
        c["fips"] = f["id"]
        ring = largest_ring(f["geometry"])
        c["lon"] = round(sum(p[0] for p in ring) / len(ring), 5)
        c["lat"] = round(sum(p[1] for p in ring) / len(ring), 5)
        out_feats.append({"type": "Feature", "id": f["id"],
                          "properties": {"name": c["name"]},
                          "geometry": f["geometry"]})
    if len(feats) != len(counties):
        raise SystemExit(f"geojson has {len(feats)} counties for state "
                         f"{state_fips} but workbook lists {len(counties)}")
    return {"type": "FeatureCollection", "features": out_feats}


def main():
    st = (sys.argv[1] if len(sys.argv) > 1 else "in").lower()
    if st not in STATE_FIPS:
        raise SystemExit(f"unknown state {st!r} — add it to STATE_FIPS/EXPECTED_N")
    chart_a = Path(sys.argv[2]) if len(sys.argv) > 2 else (
        HERE / "PDAExamples" / "Chart As" / "Validated"
        / "Floyd - Indiana (June 6 - 11, 2026).xlsx")
    us_geo = Path(sys.argv[3]) if len(sys.argv) > 3 else REPO / "counties_us.json"

    wb = openpyxl.load_workbook(chart_a, read_only=True)
    sheets = set(wb.sheetnames)
    wb.close()
    if "County Population and Threshold" in sheets:
        counties = read_populations(chart_a)          # a Chart A of this state
    elif "County Populations" in sheets:
        counties = read_populations_national(chart_a, st)   # any Table A master
        # the national sheet carries non-county rows (Grand Total, tribal
        # nations — e.g. MN's Prairie Island Indian Community). Tribal damage
        # reaches the board through county Chart A rows, so the map dim keeps
        # census counties only; every exclusion is printed, never silent.
        src = json.loads(us_geo.read_text(encoding="utf-8"))
        geonames = {_norm(f["properties"]["NAME"]) for f in src["features"]
                    if f.get("properties", {}).get("STATE") == STATE_FIPS[st]}
        keep, skipped = [], []
        for c in counties:
            key = NAME_ALIASES.get(_norm(c["name"]), _norm(c["name"]))
            (keep if key in geonames else skipped).append(c)
        for c in skipped:
            print(f"  skipped non-county row: {c['name']!r} (pop {c['population']:,})")
        counties = keep
    else:
        raise SystemExit(f"{chart_a} has neither population sheet")
    if len(counties) != EXPECTED_N[st]:
        raise SystemExit(f"expected {EXPECTED_N[st]} {st.upper()} counties, "
                         f"got {len(counties)} from {chart_a}")
    geo = build_geo(us_geo, counties, STATE_FIPS[st])

    DATA.mkdir(exist_ok=True)
    with open(DATA / f"dim_county_{st}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["fips", "name", "name_full", "population", "lat", "lon"])
        w.writeheader()
        w.writerows({k: c[k] for k in w.fieldnames} for c in counties)
    (DATA / f"{st}_counties_geo.json").write_text(json.dumps(geo), encoding="utf-8")
    print(f"dim_county_{st}.csv: {len(counties)} counties; "
          f"{st}_counties_geo.json: {len(geo['features'])} features "
          f"({(DATA / f'{st}_counties_geo.json').stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
